"""Unit tests for Text and Typography Rendering (PR11)."""

from pathlib import Path
from pptx import Presentation
from pptx.enum.text import PP_ALIGN

from backend.layout.schema import BlockRole, ElementStyle, ElementType, LayoutElement, LayoutSpec, Rect, TextStyle, VisualIntent
from backend.renderer.pptx_builder import PPTXBuilder
from backend.renderer.theme import AcademicTheme


def test_render_textbox_basic(tmp_path: Path):
    builder = PPTXBuilder()
    theme = AcademicTheme()

    slide_spec = LayoutSpec(
        slide_id="slide_1",
        slide_index=1,
        visual_intent=VisualIntent.TITLE_HERO,
    )
    builder.add_slide(slide_spec)

    text_elem = LayoutElement(
        element_id="el_title",
        element_type=ElementType.TEXT,
        role=BlockRole.HEADING,
        geometry=Rect(x=100.0, y=150.0, width=800.0, height=80.0),
        style=ElementStyle(
            text=TextStyle(
                font_size=28.0,
                font_weight="bold",
                alignment="center",
                color="#0F172A",
            )
        ),
        content="Transformer Architecture for Vision",
    )
    shape = builder.add_text(text_elem, theme)

    out_file = tmp_path / "test_text.pptx"
    builder.save(out_file)

    # Read back using python-pptx
    prs = Presentation(str(out_file))
    assert len(prs.slides) == 1
    s = prs.slides[0]
    assert len(s.shapes) == 1
    sh = s.shapes[0]
    assert sh.has_text_frame
    assert sh.text_frame.text == "Transformer Architecture for Vision"
    p = sh.text_frame.paragraphs[0]
    assert p.alignment == PP_ALIGN.CENTER
    assert p.runs[0].font.bold is True


def test_render_multiline_bullets(tmp_path: Path):
    builder = PPTXBuilder()
    theme = AcademicTheme()

    slide_spec = LayoutSpec(
        slide_id="slide_2",
        slide_index=2,
        visual_intent=VisualIntent.KEY_TAKEAWAY_LIST,
    )
    builder.add_slide(slide_spec)

    bullets = [
        "First key discovery on latency reduction",
        "Second result regarding benchmark accuracy",
        "Third ablation analysis on attention heads",
    ]
    bullet_elem = LayoutElement(
        element_id="el_bullets",
        element_type=ElementType.TEXT,
        role=BlockRole.BULLET_ITEM,
        geometry=Rect(x=80.0, y=200.0, width=600.0, height=250.0),
        style=ElementStyle(
            text=TextStyle(font_size=16.0, alignment="left")
        ),
        content=bullets,
    )
    builder.add_text(bullet_elem, theme)

    out_file = tmp_path / "test_bullets.pptx"
    builder.save(out_file)

    prs = Presentation(str(out_file))
    sh = prs.slides[0].shapes[0]
    paras = sh.text_frame.paragraphs
    assert len(paras) == 3
    assert paras[0].text == bullets[0]
    assert paras[1].text == bullets[1]
    assert paras[2].text == bullets[2]


def test_render_badge_element(tmp_path: Path):
    builder = PPTXBuilder()
    theme = AcademicTheme()

    slide_spec = LayoutSpec(
        slide_id="slide_3",
        slide_index=3,
        visual_intent=VisualIntent.TITLE_HERO,
    )
    builder.add_slide(slide_spec)

    badge_elem = LayoutElement(
        element_id="el_badge_1",
        element_type=ElementType.BADGE,
        geometry=Rect(x=100.0, y=400.0, width=150.0, height=36.0),
        style=ElementStyle(
            background_color="#DBEAFE",
            border_color="#93C5FD",
            border_width=1.0,
            corner_radius=18.0,
            text=TextStyle(font_size=11.0, font_weight="bold", color="#1E40AF"),
        ),
        content="CVPR 2024 Oral",
    )
    builder.add_shape(badge_elem, theme)

    out_file = tmp_path / "test_badge.pptx"
    builder.save(out_file)

    prs = Presentation(str(out_file))
    sh = prs.slides[0].shapes[0]
    assert sh.has_text_frame
    assert sh.text_frame.text == "CVPR 2024 Oral"


def test_speaker_notes_rendering(tmp_path: Path):
    builder = PPTXBuilder()
    slide_spec = LayoutSpec(
        slide_id="slide_notes",
        slide_index=1,
        visual_intent=VisualIntent.TITLE_HERO,
        speaker_notes="Welcome the audience and introduce the paper's key claim.",
    )
    builder.add_slide(slide_spec)

    out_file = tmp_path / "test_notes.pptx"
    builder.save(out_file)

    prs = Presentation(str(out_file))
    notes = prs.slides[0].notes_slide.notes_text_frame.text
    assert "Welcome the audience" in notes
