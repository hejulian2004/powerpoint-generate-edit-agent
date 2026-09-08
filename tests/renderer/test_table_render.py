"""Unit tests for Table Rendering (PR11)."""

from pathlib import Path
from pptx import Presentation

from backend.layout.schema import ElementType, LayoutElement, LayoutSpec, Rect, VisualIntent
from backend.renderer.assets import AssetResolver
from backend.renderer.pptx_builder import PPTXBuilder
from backend.renderer.theme import AcademicTheme


def test_render_table_basic(tmp_path: Path):
    builder = PPTXBuilder()
    theme = AcademicTheme()

    slide_spec = LayoutSpec(
        slide_id="slide_tbl",
        slide_index=1,
        visual_intent=VisualIntent.BENCHMARK_COMPARISON,
    )
    builder.add_slide(slide_spec)

    table_elem = LayoutElement(
        element_id="el_tbl_1",
        element_type=ElementType.TABLE,
        geometry=Rect(x=100.0, y=150.0, width=700.0, height=350.0),
        content={"source_table_id": "tbl_benchmark"},
    )
    table_data = {
        "header": ["Method", "Accuracy", "Latency (ms)"],
        "rows": [
            ["ResNet-50", "76.2%", "15.4"],
            ["ViT-B/16", "81.8%", "32.1"],
            ["Ours (FastViT)", "84.5%", "12.0"],
        ],
        "highlight_cells": ["2,0", "2,1", "2,2"],
    }
    builder.add_table(table_elem, table_data, theme)

    out_file = tmp_path / "test_table_render.pptx"
    builder.save(out_file)

    prs = Presentation(str(out_file))
    slide = prs.slides[0]
    assert len(slide.shapes) == 1
    tbl_shape = slide.shapes[0]
    assert tbl_shape.has_table
    table = tbl_shape.table

    # 1 header + 3 rows = 4 rows, 3 cols
    assert len(table.rows) == 4
    assert len(table.columns) == 3

    # Check header text
    assert table.cell(0, 0).text_frame.text == "Method"
    assert table.cell(0, 1).text_frame.text == "Accuracy"

    # Check cell text
    assert table.cell(1, 0).text_frame.text == "ResNet-50"
    assert table.cell(3, 0).text_frame.text == "Ours (FastViT)"


def test_table_resolver_fallback(tmp_path: Path):
    resolver = AssetResolver()
    resolved = resolver.resolve_table("table1")
    assert "header" in resolved
    assert "rows" in resolved
    assert len(resolved["rows"]) >= 2
    assert len(resolved["header"]) >= 2
