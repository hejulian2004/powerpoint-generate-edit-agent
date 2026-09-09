"""Unit tests for Figure and Image Asset Rendering (PR11)."""

from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from backend.layout.schema import BlockRole, ElementStyle, ElementType, LayoutElement, LayoutSpec, Rect, VisualIntent
from backend.renderer.assets import AssetResolver
from backend.renderer.pptx_builder import PPTXBuilder
from backend.renderer.theme import AcademicTheme


def test_render_figure_from_real_file(tmp_path: Path):
    # 1. Create a dummy image
    img_path = tmp_path / "real_fig.png"
    img = Image.new("RGB", (400, 300), color=(100, 150, 200))
    img.save(img_path)

    builder = PPTXBuilder()
    theme = AcademicTheme()

    slide_spec = LayoutSpec(
        slide_id="slide_fig",
        slide_index=1,
        visual_intent=VisualIntent.PIPELINE_ARCHITECTURE,
    )
    builder.add_slide(slide_spec)

    fig_elem = LayoutElement(
        element_id="el_fig_1",
        element_type=ElementType.FIGURE,
        role=BlockRole.CALLOUT,
        geometry=Rect(x=600.0, y=120.0, width=500.0, height=400.0),
        content={"source_figure_id": "fig_pipeline"},
    )
    builder.add_image(fig_elem, img_path, theme)

    out_file = tmp_path / "test_fig_render.pptx"
    builder.save(out_file)

    prs = Presentation(str(out_file))
    slide = prs.slides[0]
    assert len(slide.shapes) == 1
    sh = slide.shapes[0]
    assert sh.shape_type == MSO_SHAPE_TYPE.PICTURE


def test_asset_resolver_missing_figure_raises(tmp_path: Path):
    import pytest
    resolver = AssetResolver(cache_dir=tmp_path / "cache")
    with pytest.raises(FileNotFoundError, match="could not be resolved"):
        resolver.resolve_figure(
            figure_id="fig_unknown_model",
            caption_hint="Proposed Multi-Modal Transformer Architecture",
        )


def test_render_figure_with_resolver(tmp_path: Path):
    assets_dir = tmp_path / "assets"
    assets_dir.mkdir()
    img_path = assets_dir / "fig_auto_placeholder.png"
    Image.new("RGB", (400, 300), color=(100, 150, 200)).save(img_path)

    resolver = AssetResolver(assets_dir=assets_dir)
    builder = PPTXBuilder()
    theme = AcademicTheme()

    slide_spec = LayoutSpec(
        slide_id="slide_res",
        slide_index=1,
        visual_intent=VisualIntent.PIPELINE_ARCHITECTURE,
    )
    builder.add_slide(slide_spec)

    fig_elem = LayoutElement(
        element_id="el_fig_synth",
        element_type=ElementType.FIGURE,
        geometry=Rect(x=100.0, y=100.0, width=400.0, height=300.0),
        content={"source_figure_id": "fig_auto_placeholder", "caption": "Framework Overview"},
    )
    resolved_img = resolver.resolve_figure(fig_elem.content["source_figure_id"], caption_hint=fig_elem.content["caption"])
    builder.add_image(fig_elem, resolved_img, theme)

    out_file = tmp_path / "test_fig_synth.pptx"
    builder.save(out_file)

    prs = Presentation(str(out_file))
    sh = prs.slides[0].shapes[0]
    assert sh.shape_type == MSO_SHAPE_TYPE.PICTURE


def test_asset_resolver_special_chars_and_existing_extension(tmp_path: Path):
    # 1. Existing file with extension inside assets_dir
    assets_dir = tmp_path / "assets"
    assets_dir.mkdir()
    existing_img = assets_dir / "diagram.png"
    Image.new("RGB", (200, 200), color=(50, 100, 150)).save(existing_img)

    resolver = AssetResolver(assets_dir=assets_dir, cache_dir=tmp_path / "cache")
    res_direct = resolver.resolve_figure("diagram.png")
    assert res_direct == existing_img

    # 2. Missing asset raises FileNotFoundError
    import pytest
    with pytest.raises(FileNotFoundError, match="could not be resolved"):
        resolver.resolve_figure("nonexistent_fig", allow_synthetic=False)
