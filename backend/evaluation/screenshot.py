"""Cross-platform PPTX Slide Screenshot Renderer Pipeline (PR12).

Provides high-fidelity, deterministic slide-to-image rasterization supporting:
- Windows Native (PowerPoint COM automation via PowerShell)
- Linux / Cross-platform (LibreOffice headless + pypdfium2 PDF rendering)
- Headless Fallback (python-pptx + Pillow deterministic rasterization)
"""

from __future__ import annotations

import hashlib
import os
import shutil
import subprocess
from abc import ABC, abstractmethod
from pathlib import Path
from typing import List, Optional, Tuple, Union

from PIL import Image, ImageDraw, ImageFont


class ScreenshotBackend(ABC):
    """Abstract base class for PPTX slide rendering backends."""

    @abstractmethod
    def is_available(self) -> bool:
        """Check whether runtime dependencies for this backend are present."""
        pass

    @abstractmethod
    def render(
        self,
        pptx_path: Path,
        output_dir: Path,
        resolution: Tuple[int, int] = (1280, 720),
    ) -> List[Path]:
        """Render all slides in the presentation to PNG images.

        Returns sorted list of output PNG paths: [1.png, 2.png, ...].
        """
        pass


class PowerPointBackend(ScreenshotBackend):
    """Native Windows PowerPoint COM automation backend.

    Guarantees native Office typography and fidelity when run on Windows.
    """

    def is_available(self) -> bool:
        if os.name != "nt":
            return False
        # Check if PowerShell is available
        return shutil.which("powershell") is not None

    def render(
        self,
        pptx_path: Path,
        output_dir: Path,
        resolution: Tuple[int, int] = (1280, 720),
    ) -> List[Path]:
        in_path = pptx_path.resolve()
        out_dir = output_dir.resolve()
        out_dir.mkdir(parents=True, exist_ok=True)

        safe_in_path = str(in_path).replace("'", "''")
        safe_out_dir = str(out_dir).replace("'", "''")
        width_px, height_px = resolution

        ps_script = f"""
$ErrorActionPreference = 'Stop'
$ppt = $null
$pres = $null
try {{
    $ppt = New-Object -ComObject PowerPoint.Application
    $pres = $ppt.Presentations.Open('{safe_in_path}', [Microsoft.Office.Core.MsoTriState]::msoTrue, [Microsoft.Office.Core.MsoTriState]::msoFalse, [Microsoft.Office.Core.MsoTriState]::msoFalse)
    for ($i = 1; $i -le $pres.Slides.Count; $i++) {{
        $slidePath = Join-Path '{safe_out_dir}' ("$i.png")
        $pres.Slides.Item($i).Export($slidePath, "PNG", {width_px}, {height_px})
    }}
    Write-Host "COM_EXPORT_SUCCESS"
}} catch {{
    Write-Host ("COM_EXPORT_ERROR: " + $_.Exception.Message)
}} finally {{
    if ($pres) {{
        try {{ $pres.Close() }} catch {{}}
        try {{ [System.Runtime.Interopservices.Marshal]::ReleaseComObject($pres) | Out-Null }} catch {{}}
    }}
    if ($ppt) {{
        try {{ $ppt.Quit() }} catch {{}}
        try {{ [System.Runtime.Interopservices.Marshal]::ReleaseComObject($ppt) | Out-Null }} catch {{}}
    }}
}}
"""
        try:
            proc = subprocess.run(
                ["powershell", "-NoProfile", "-Command", ps_script],
                capture_output=True,
                text=True,
                timeout=60,
            )
            if "COM_EXPORT_SUCCESS" in proc.stdout:
                pngs = self._collect_sorted_pngs(out_dir)
                if pngs:
                    return pngs
        except Exception:
            pass

        return []

    @staticmethod
    def _collect_sorted_pngs(directory: Path) -> List[Path]:
        candidates = list(directory.glob("*.png"))

        def _sort_key(p: Path) -> int:
            name = p.stem.replace("slide_", "")
            return int(name) if name.isdigit() else 999999

        return sorted([p for p in candidates if p.stem.replace("slide_", "").isdigit()], key=_sort_key)


class LibreOfficeBackend(ScreenshotBackend):
    """Linux/cross-platform headless LibreOffice + pypdfium2 renderer."""

    def __init__(self) -> None:
        self.soffice_cmd = shutil.which("soffice") or shutil.which("libreoffice")

    def is_available(self) -> bool:
        return self.soffice_cmd is not None

    def render(
        self,
        pptx_path: Path,
        output_dir: Path,
        resolution: Tuple[int, int] = (1280, 720),
    ) -> List[Path]:
        if not self.soffice_cmd:
            return []

        in_path = pptx_path.resolve()
        out_dir = output_dir.resolve()
        out_dir.mkdir(parents=True, exist_ok=True)
        pdf_dir = out_dir / "temp_pdf_render"
        pdf_dir.mkdir(parents=True, exist_ok=True)

        try:
            subprocess.run(
                [self.soffice_cmd, "--headless", "--convert-to", "pdf", str(in_path), "--outdir", str(pdf_dir)],
                check=True,
                capture_output=True,
                timeout=60,
            )
            pdf_file = pdf_dir / f"{in_path.stem}.pdf"
            if not pdf_file.is_file():
                return []

            import pypdfium2 as pdfium

            pdf = pdfium.PdfDocument(pdf_file)
            target_w, target_h = resolution
            rendered: List[Path] = []

            for page_idx in range(len(pdf)):
                page = pdf[page_idx]
                # Render to PIL Image
                bitmap = page.render(scale=2.0)
                img = bitmap.to_pil()
                if img.size != (target_w, target_h):
                    img = img.resize((target_w, target_h), Image.Resampling.LANCZOS)
                dest = out_dir / f"{page_idx + 1}.png"
                img.save(dest, format="PNG")
                rendered.append(dest)

            # Cleanup temp pdf
            try:
                shutil.rmtree(pdf_dir, ignore_errors=True)
            except Exception:
                pass

            return rendered
        except Exception:
            return []


class FallbackScreenshotBackend(ScreenshotBackend):
    """Deterministic fallback rasterizer when neither PowerPoint nor LibreOffice is present.

    Reads the presentation via python-pptx and generates deterministic slide layouts
    using Pillow. Ensures tests and CI pipelines run without external software dependencies.
    """

    def is_available(self) -> bool:
        return True

    def render(
        self,
        pptx_path: Path,
        output_dir: Path,
        resolution: Tuple[int, int] = (1280, 720),
    ) -> List[Path]:
        import pptx
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        prs = pptx.Presentation(str(pptx_path))
        out_dir = output_dir.resolve()
        out_dir.mkdir(parents=True, exist_ok=True)

        width_px, height_px = resolution
        slide_width_emu = prs.slide_width
        slide_height_emu = prs.slide_height
        scale_x = width_px / float(slide_width_emu)
        scale_y = height_px / float(slide_height_emu)

        rendered: List[Path] = []

        for slide_idx, slide in enumerate(prs.slides, start=1):
            img = Image.new("RGB", (width_px, height_px), color=(255, 255, 255))
            draw = ImageDraw.Draw(img)

            # Draw background
            draw.rectangle([(0, 0), (width_px - 1, height_px - 1)], fill=(248, 250, 252))

            for shape in slide.shapes:
                sx = int(shape.left * scale_x)
                sy = int(shape.top * scale_y)
                sw = int(shape.width * scale_x)
                sh = int(shape.height * scale_y)

                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    # Draw picture placeholder box
                    draw.rectangle([(sx, sy), (sx + sw, sy + sh)], fill=(226, 232, 240), outline=(148, 163, 184), width=2)
                    draw.text((sx + 8, sy + 8), "[Picture Asset]", fill=(71, 85, 105))
                elif shape.has_table:
                    # Draw table placeholder
                    draw.rectangle([(sx, sy), (sx + sw, sy + sh)], fill=(255, 255, 255), outline=(59, 130, 246), width=2)
                    draw.text((sx + 8, sy + 8), f"[Table: {len(shape.table.rows)}x{len(shape.table.columns)}]", fill=(30, 58, 138))
                elif shape.has_text_frame:
                    # Draw text container and text
                    draw.rectangle([(sx, sy), (sx + sw, sy + sh)], outline=(229, 231, 235), width=1)
                    text_str = shape.text_frame.text[:120].strip()
                    if text_str:
                        draw.text((sx + 6, sy + 6), text_str, fill=(15, 23, 42))
                else:
                    # Generic shape
                    draw.rectangle([(sx, sy), (sx + sw, sy + sh)], fill=(241, 245, 249), outline=(203, 213, 225), width=1)

            dest_path = out_dir / f"{slide_idx}.png"
            img.save(dest_path, format="PNG")
            rendered.append(dest_path)

        return rendered


class ScreenshotRenderer:
    """Orchestrator for exporting PPTX presentations into slide screenshot PNGs."""

    def __init__(
        self,
        backend: Optional[ScreenshotBackend] = None,
        force_fallback: bool = False,
    ) -> None:
        self.force_fallback = force_fallback
        if backend is not None:
            self.backend = backend
        elif force_fallback:
            self.backend = FallbackScreenshotBackend()
        else:
            self.backend = self._auto_select_backend()

    @staticmethod
    def _auto_select_backend() -> ScreenshotBackend:
        ppt_be = PowerPointBackend()
        if ppt_be.is_available():
            return ppt_be
        lo_be = LibreOfficeBackend()
        if lo_be.is_available():
            return lo_be
        return FallbackScreenshotBackend()

    def render_screenshots(
        self,
        pptx_path: Union[str, Path],
        output_dir: Union[str, Path],
        resolution: Tuple[int, int] = (1280, 720),
    ) -> List[Path]:
        """Render presentation slides to deterministic PNG images."""
        p_path = Path(pptx_path).resolve()
        o_dir = Path(output_dir).resolve()
        if not p_path.is_file():
            raise FileNotFoundError(f"PPTX file not found: {p_path}")

        o_dir.mkdir(parents=True, exist_ok=True)
        results = self.backend.render(p_path, o_dir, resolution=resolution)

        # If primary backend yielded nothing, gracefully retry with fallback
        if not results and not isinstance(self.backend, FallbackScreenshotBackend):
            fallback = FallbackScreenshotBackend()
            results = fallback.render(p_path, o_dir, resolution=resolution)

        return results

    @staticmethod
    def compute_image_hash(image_path: Union[str, Path]) -> str:
        """Compute SHA-256 checksum of an exported screenshot."""
        data = Path(image_path).read_bytes()
        return hashlib.sha256(data).hexdigest()


def render_screenshots(
    pptx_path: Union[str, Path],
    output_dir: Union[str, Path],
    resolution: Tuple[int, int] = (1280, 720),
) -> List[Path]:
    """Convenience functional interface for rendering PPTX slides to PNG."""
    renderer = ScreenshotRenderer()
    return renderer.render_screenshots(pptx_path, output_dir, resolution=resolution)
